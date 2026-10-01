#!/usr/bin/env python3
"""Rigenera e verifica la vista XLSX del registro VC compatto.

Il CSV resta l'unica fonte operativa. La vista Excel mostra collegamenti veri
nelle colonne ``apri_md`` ed ``evidenza_vc``, blocca prima riga e prima
colonna e colora lo stato VC. Non va modificata a mano.

Uso:

    python strumenti/genera-vista-xlsx-registro.py
    python strumenti/genera-vista-xlsx-registro.py --verifica-soltanto
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
CSV_IN = ROOT / "registro-verifica-conformita-ntc2018.csv"
XLSX_OUT = ROOT / "registro-verifica-conformita-ntc2018.xlsx"
FOGLIO = "registro-vc"

FIELDS = [
    "paragrafo",
    "titolo",
    "n.",
    "apri_md",
    "pagine_pdf_ntc",
    "pagine_pdf_circolare",
    "stato_vc",
    "data_ultima_vc",
    "prossima_azione",
    "evidenza_vc",
    "note",
    "procedura_vc",
    "agente_vc",
]
LINK_FIELDS = {"apri_md", "evidenza_vc"}
STATES = ["VC-OK", "IN-CORSO", "DA-RICONFERMARE", "DA-RIFARE", "MAI-FATTA"]
WIDTHS = {
    "paragrafo": 15,
    "titolo": 44,
    "n.": 7,
    "apri_md": 11,
    "pagine_pdf_ntc": 17,
    "pagine_pdf_circolare": 23,
    "stato_vc": 20,
    "data_ultima_vc": 17,
    "prossima_azione": 48,
    "evidenza_vc": 13,
    "note": 48,
    "procedura_vc": 32,
    "agente_vc": 24,
}
STATE_COLORS = {
    "VC-OK": "C6EFCE",
    "IN-CORSO": "BDD7EE",
    "DA-RICONFERMARE": "FFF2CC",
    "DA-RIFARE": "F4CCCC",
    "MAI-FATTA": "E7E6E6",
}


def read_csv() -> list[dict[str, str]]:
    with CSV_IN.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        if reader.fieldnames not in (FIELDS, FIELDS[:-1], FIELDS[:-2]):
            raise SystemExit(
                "intestazione CSV inattesa:\n"
                f"  ottenuta: {reader.fieldnames}\n"
                f"  attesa  : {FIELDS}"
            )
        rows = list(reader)
    if any(None in row for row in rows):
        raise SystemExit("CSV con righe di larghezza incoerente")
    for row in rows:
        row.setdefault("procedura_vc", "")
        row.setdefault("agente_vc", "")
    return rows


def add_link(cell, target: str, label: str = "Apri") -> None:
    cell.value = label
    cell.hyperlink = target
    cell.style = "Hyperlink"
    cell.alignment = Alignment(horizontal="center", vertical="center")


def generate(rows: list[dict[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = FOGLIO
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = "B2"
    sheet.append(FIELDS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_side = Side(style="thin", color="9EADBA")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=header_side)
    sheet.row_dimensions[1].height = 32

    for row_index, row in enumerate(rows, start=2):
        sheet.append([None] * len(FIELDS))
        for column_index, field in enumerate(FIELDS, start=1):
            value = row[field]
            cell = sheet.cell(row=row_index, column=column_index)
            if field in LINK_FIELDS and value:
                target = (ROOT / Path(value)).resolve()
                if not target.is_file():
                    raise SystemExit(f"collegamento inesistente alla riga {row_index}: {value}")
                add_link(cell, value)
            elif not value:
                cell.value = None
            elif field == "n.":
                cell.value = int(value)
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.value = value

        sheet.cell(row=row_index, column=1).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        for field in (
            "pagine_pdf_ntc",
            "pagine_pdf_circolare",
            "stato_vc",
            "data_ultima_vc",
        ):
            column = FIELDS.index(field) + 1
            sheet.cell(row=row_index, column=column).alignment = Alignment(
                horizontal="center", vertical="top", wrap_text=True
            )
        for field in ("titolo", "prossima_azione", "note", "procedura_vc", "agente_vc"):
            column = FIELDS.index(field) + 1
            sheet.cell(row=row_index, column=column).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    last_row = len(rows) + 1
    last_column = get_column_letter(len(FIELDS))
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    for index, field in enumerate(FIELDS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[field]

    state_column = get_column_letter(FIELDS.index("stato_vc") + 1)
    state_range = f"{state_column}2:{state_column}{last_row}"
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(STATES) + '"',
        allow_blank=False,
    )
    validation.error = "Usare uno degli stati VC previsti."
    validation.errorTitle = "Stato VC non valido"
    validation.prompt = "Scegliere lo stato VC dall'elenco."
    validation.promptTitle = "Stato VC"
    sheet.add_data_validation(validation)
    validation.add(state_range)

    for state, color in STATE_COLORS.items():
        rule = FormulaRule(
            formula=[f'${state_column}2="{state}"'],
            fill=PatternFill("solid", fgColor=color),
        )
        sheet.conditional_formatting.add(state_range, rule)

    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "Pagina &P di &N"
    sheet.oddFooter.right.text = "Vista derivata dal CSV"

    workbook.save(XLSX_OUT)
    print(f"Scritto: {XLSX_OUT}")
    print(f"Righe dati: {len(rows)}; colonne: {len(FIELDS)}")
    print(f"Riquadro congelato: {sheet.freeze_panes}")


def verify(rows: list[dict[str, str]]) -> int:
    workbook = load_workbook(XLSX_OUT, data_only=False)
    errors: list[str] = []
    if workbook.sheetnames != [FOGLIO]:
        errors.append(f"fogli inattesi: {workbook.sheetnames}")
    sheet = workbook[FOGLIO]
    if sheet.max_row != len(rows) + 1 or sheet.max_column != len(FIELDS):
        errors.append(
            f"dimensioni {sheet.max_row - 1}x{sheet.max_column}, "
            f"attese {len(rows)}x{len(FIELDS)}"
        )
    if [cell.value for cell in sheet[1]] != FIELDS:
        errors.append("intestazione diversa dal CSV")
    if sheet.freeze_panes != "B2":
        errors.append(f"riquadro congelato {sheet.freeze_panes}, atteso B2")
    expected_filter = f"A1:{get_column_letter(len(FIELDS))}{len(rows) + 1}"
    if sheet.auto_filter.ref != expected_filter:
        errors.append(f"filtro {sheet.auto_filter.ref}, atteso {expected_filter}")

    links = 0
    formula_cells: list[str] = []
    for row_index, expected in enumerate(rows, start=2):
        for column_index, field in enumerate(FIELDS, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            wanted = expected[field]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells.append(cell.coordinate)
            if field in LINK_FIELDS and wanted:
                links += 1
                target = cell.hyperlink.target if cell.hyperlink else ""
                if cell.value != "Apri" or target != wanted:
                    errors.append(
                        f"{cell.coordinate}: link ottenuto value={cell.value!r}, "
                        f"target={target!r}; atteso {wanted!r}"
                    )
                continue
            obtained = "" if cell.value is None else str(cell.value)
            if obtained != wanted:
                errors.append(
                    f"{cell.coordinate} ({field}): XLSX={obtained!r}, CSV={wanted!r}"
                )

    if formula_cells:
        errors.append(f"formule inattese nelle celle: {', '.join(formula_cells[:20])}")
    if len(sheet.data_validations.dataValidation) != 1:
        errors.append("convalida dati dello stato VC assente o duplicata")
    color_rules = sum(
        len(rules) for rules in sheet.conditional_formatting._cf_rules.values()
    )
    if color_rules != len(STATE_COLORS):
        errors.append("regole colore degli stati VC incomplete")

    print(f"Verifica: {len(rows)} righe, {len(FIELDS)} colonne, {links} collegamenti")
    print(f"Differenze non dichiarate: {len(errors)}")
    for error in errors[:50]:
        print(f"KO: {error}")
    if errors:
        if len(errors) > 50:
            print(f"... e altri {len(errors) - 50} errori")
        return 1
    print("ESITO: vista conforme al CSV")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--verifica-soltanto",
        action="store_true",
        help="non riscrive la vista, la confronta soltanto con il CSV",
    )
    arguments = parser.parse_args()
    if not CSV_IN.is_file():
        print(f"CSV non trovato: {CSV_IN}", file=sys.stderr)
        return 1
    rows = read_csv()
    if arguments.verifica_soltanto:
        if not XLSX_OUT.is_file():
            print(f"vista non trovata: {XLSX_OUT}", file=sys.stderr)
            return 1
    else:
        generate(rows)
    return verify(rows)


if __name__ == "__main__":
    raise SystemExit(main())
